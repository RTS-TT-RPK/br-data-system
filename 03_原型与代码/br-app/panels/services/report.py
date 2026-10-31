"""一键 Excel 报表：参照《综合物料报表-修订版》的配色/表头/汇总行风格。"""
from __future__ import annotations

import io
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter

from panels.models import Component, PanelDefinition, Record
from panels.services.analysis import component_summary, material_component_summary, summarize_by_material
from panels.services.component_parser import component_sum_percent

NAVY = "1B2A4A"
LIGHT_BLUE = "D6E4F0"
BORDER_COLOR = "B7C5D9"
TEXT = "37352F"

TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color=NAVY)
INFO_FONT = Font(name="微软雅黑", size=10, bold=True, color=NAVY)
HEADER_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
SECTION_FONT = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
DATA_FONT = Font(name="微软雅黑", size=10.5, color=TEXT)
TOTAL_FONT = Font(name="微软雅黑", size=12, bold=True, color=NAVY)
SMALL_FONT = Font(name="微软雅黑", size=9, color=TEXT)

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
LIGHT_FILL = PatternFill("solid", fgColor=LIGHT_BLUE)
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
THIN = Side(style="thin", color=BORDER_COLOR)
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _title(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = TITLE_FONT
    c.alignment = CENTER
    ws.row_dimensions[row].height = 30
    return row + 1


def _info(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = INFO_FONT
    c.fill = LIGHT_FILL
    c.alignment = CENTER
    c.border = BORDER
    ws.row_dimensions[row].height = 20
    return row + 1


def _section(ws, row, ncols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = SECTION_FONT
    c.fill = HEADER_FILL
    c.alignment = LEFT
    c.border = BORDER
    ws.row_dimensions[row].height = 24
    return row + 1


def _header(ws, row, headers):
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = 24
    return row + 1


def _row(ws, row, values, fonts=None, fills=None, aligns=None):
    for i, value in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=value)
        c.font = (fonts or DATA_FONT)
        c.border = BORDER
        c.alignment = (aligns or CENTER)
        if fills:
            c.fill = fills
    return row + 1


def _print_setup(ws, landscape=True, fit_width=1):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _pct(v):
    if v is None:
        return ""
    return round(float(v), 2)


def _water_pct(record):
    vals = [float(c.numeric_value) * 100 for c in record.components.all() if c.category == "water" and c.unit == "%" and c.numeric_value is not None]
    return max(vals) if vals else None


def _inventory_overview():
    groups = defaultdict(list)
    for rec in Record.objects.filter(panel_key="br_inventory", status=Record.STATUS_CONFIRMED).prefetch_related("components"):
        name = (rec.data_json or {}).get("名称") or "（未填）"
        groups[name].append(rec)
    rows = []
    for name, recs in groups.items():
        if name == "（未填）":
            continue
        waters = []
        for r in recs:
            w = _water_pct(r)
            if w is not None:
                waters.append(w)
        sums = [component_sum_percent(list(r.components.all())) for r in recs]
        rows.append(
            {
                "material": name,
                "count": len(recs),
                "avg_water": sum(waters) / len(waters) if waters else None,
                "max_water": max(waters) if waters else None,
                "avg_sum": sum(sums) / len(sums) if sums else None,
            }
        )
    rows.sort(key=lambda x: (-x["count"], x["material"]))
    return rows


def build_summary_report() -> io.BytesIO:
    from panels.services.analysis import analyze_incoming
    analysis = analyze_incoming()
    by_material = summarize_by_material(analysis)
    comps = component_summary()
    mat_comps = material_component_summary()
    inventory = _inventory_overview()
    inv_records = list(
        Record.objects.filter(panel_key="br_inventory", status=Record.STATUS_CONFIRMED).order_by("id")
    )
    hw_records = list(Record.objects.filter(panel_key="br_hardware", status=Record.STATUS_CONFIRMED).order_by("id"))
    inv_panel = PanelDefinition.objects.filter(key="br_inventory").first()
    hw_panel = PanelDefinition.objects.filter(key="br_hardware").first()

    wb = Workbook()
    ws = wb.active
    ws.title = "综合汇总"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY

    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    ncols = 10
    row = _title(ws, 1, ncols, "示例企业物料综合汇总表")
    row = _info(ws, row, ncols, f"单位：吨 / 百分比　　统计时间：{now}　　数据来源：面板数据预览")
    row += 1

    # KPI 卡片
    kpis = [
        ("物料种类", len({(r.data_json or {}).get('中文名') for r in Record.objects.filter(panel_key='br_material_property')})),
        ("库存记录", len(inv_records)),
        ("来料分析", len(analysis)),
        ("成分记录", Component.objects.count()),
        ("五金配件", len(hw_records)),
    ]
    kpi_cols = [1, 3, 5, 7, 9]
    for (label, value), col in zip(kpis, kpi_cols):
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        c1 = ws.cell(row=row, column=col, value=label)
        c1.font = INFO_FONT
        c1.fill = LIGHT_FILL
        c1.alignment = CENTER
        c1.border = BORDER
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        c2 = ws.cell(row=row + 1, column=col, value=value)
        c2.font = Font(name="微软雅黑", size=18, bold=True, color=NAVY)
        c2.alignment = CENTER
        c2.border = BORDER
    row += 3

    # 来料分析按物料
    row = _section(ws, row, ncols, "来料分析 · 按物料汇总")
    row = _header(ws, row, ["物料", "记录数", "主溶剂（最高识别含量）", "平均含量%", "主成分沸点℃", "与其它最小沸点差℃", "共沸记录（仅二元表）", "说明"])
    for m in by_material:
        row = _row(ws, row, [
            m["material"], m["count"], m["main_name"], _pct(m["avg_pct"]), _pct(m["bp"]),
            _pct(m["avg_gap"]), m.get("azeo_label", m["azeo_flag"]), "",
        ])
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    note = ws.cell(row=row, column=1, value="口径提示：主溶剂按最高识别含量判定；共沸仅查台账二元配对，未标注压力/组成适用条件，未计算多组分汽液平衡；未收录不等于不存在共沸。")
    note.font = SMALL_FONT
    note.alignment = LEFT
    note.fill = LIGHT_FILL
    note.border = BORDER
    ws.row_dimensions[row].height = 30
    row += 2

    # 库存概览
    row = _section(ws, row, ncols, "库存概览 · 按物料")
    row = _header(ws, row, ["物料", "记录数", "平均含水%", "最高含水%", "平均成分合计%"])
    for m in inventory:
        row = _row(ws, row, [m["material"], m["count"], _pct(m["avg_water"]), _pct(m["max_water"]), _pct(m["avg_sum"])])
    row += 1

    # 五金库
    if hw_panel:
        row = _section(ws, row, ncols, "五金配件库")
        headers = [c.get("label") for c in (hw_panel.columns or [])][:ncols]
        row = _header(ws, row, headers)
        if not hw_records:
            row = _row(ws, row, ["（暂无数据）"] + [""] * (len(headers) - 1))
        else:
            for rec in hw_records:
                vals = [(rec.data_json or {}).get(c.get("key"), "") for c in (hw_panel.columns or [])][:ncols]
                row = _row(ws, row, vals)

    _widths(ws, [20, 14, 16, 12, 12, 16, 12, 28, 10, 10])
    for r in range(1, row):
        ws.row_dimensions[r].height = ws.row_dimensions[r].height or 20

    # 明细表
    ws2 = wb.create_sheet("来料分析明细")
    ws2.sheet_view.showGridLines = False
    headers = ["日期", "位置", "物料", "成分明细原文", "主溶剂（最高识别含量）", "含量%", "沸点℃", "其它成分（含量 / 沸点 / 差距）", "最小沸点差℃", "最大沸点差℃", "共沸记录（仅二元表）", "共沸说明"]
    row2 = _title(ws2, 1, len(headers), "来料分析明细")
    row2 = _info(ws2, row2, len(headers), f"共 {len(analysis)} 条　统计时间：{now}　共沸为台账二元配对查询，未计算多组分汽液平衡")
    row2 = _header(ws2, row2, headers)
    for item in analysis:
        row2 = _row(ws2, row2, [
            item["date"], item["location"], item["material"], item["raw_text"],
            item["main_name"], item["main_pct_label"], item["main_bp_label"],
            item["other_text"], item["min_gap_label"], item["max_gap_label"],
            item.get("azeo_label", item["azeo_flag"]), item["azeo_note"],
        ], aligns=LEFT)
    ws2.freeze_panes = "A4"
    ws2.auto_filter.ref = f"A3:L{row2 - 1}"
    _widths(ws2, [12, 14, 16, 46, 16, 10, 10, 60, 12, 12, 10, 36])

    # 库存明细
    ws3 = wb.create_sheet("库存明细")
    ws3.sheet_view.showGridLines = False
    inv_cols = [c.get("label") for c in (inv_panel.columns or [])]
    row3 = _title(ws3, 1, len(inv_cols), "物料库存明细")
    row3 = _info(ws3, row3, len(inv_cols), f"共 {len(inv_records)} 条")
    row3 = _header(ws3, row3, inv_cols)
    for rec in inv_records:
        vals = [(rec.data_json or {}).get(c.get("key"), "") for c in (inv_panel.columns or [])]
        row3 = _row(ws3, row3, vals, aligns=LEFT)
    ws3.freeze_panes = "A4"
    ws3.auto_filter.ref = f"A3:{get_column_letter(len(inv_cols))}{row3 - 1}"
    _widths(ws3, [14, 14, 12, 10, 16, 10, 44, 24, 12, 10, 10, 12])

    # 成分汇总
    ws4 = wb.create_sheet("成分汇总")
    ws4.sheet_view.showGridLines = False
    row4 = _title(ws4, 1, 5, "来料成分汇总")
    row4 = _info(ws4, row4, 5, "按来料分析拆出的成分统计")
    row4 = _header(ws4, row4, ["成分", "出现次数", "平均含量%", "最高%", "最低%"])
    for c in comps:
        row4 = _row(ws4, row4, [c["name"], c["count"], round(c["avg"], 3), round(c["max"], 3), round(c["min"], 3)])
    _widths(ws4, [22, 12, 14, 12, 12])

    # 物料 × 成分明细
    ws6 = wb.create_sheet("物料成分明细")
    ws6.sheet_view.showGridLines = False
    row6 = _title(ws6, 1, 6, "物料成分明细（每种物料有哪些成分、含量多少）")
    row6 = _info(ws6, row6, 6, f"按来料分析统计　共 {len(mat_comps)} 行")
    row6 = _header(ws6, row6, ["物料", "成分", "出现次数", "平均含量%", "最高%", "最低%"])
    for m in mat_comps:
        row6 = _row(ws6, row6, [m["material"], m["component"], m["count"], round(m["avg"], 3), round(m["max"], 3), round(m["min"], 3)])
    ws6.freeze_panes = "A4"
    ws6.auto_filter.ref = f"A3:F{row6 - 1}"
    _widths(ws6, [20, 20, 10, 14, 10, 10])

    # 图表
    ws5 = wb.create_sheet("图表")
    ws5.sheet_view.showGridLines = False
    ws5.cell(row=1, column=1, value="来料成分出现次数 Top 10").font = TITLE_FONT
    top = comps[:10]
    ws5.cell(row=3, column=1, value="成分")
    ws5.cell(row=3, column=2, value="出现次数")
    for i, c in enumerate(top, start=4):
        ws5.cell(row=i, column=1, value=c["name"]).font = SMALL_FONT
        ws5.cell(row=i, column=2, value=c["count"]).font = SMALL_FONT
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "来料成分出现次数 Top 10"
    data = Reference(ws5, min_col=2, min_row=3, max_row=3 + len(top))
    cats = Reference(ws5, min_col=1, min_row=4, max_row=3 + len(top))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 9
    chart.width = 18
    ws5.add_chart(chart, "D3")
    _widths(ws5, [20, 12])

    _print_setup(ws)
    _print_setup(ws2)
    _print_setup(ws3)
    _print_setup(ws4)
    _print_setup(ws6)
    _print_setup(ws5)
    ws.print_area = f"A1:J{max(1, row - 1)}"
    ws2.print_area = f"A1:L{max(1, row2 - 1)}"
    ws3.print_area = f"A1:{get_column_letter(max(1, len(inv_cols)))}{max(1, row3 - 1)}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_query_export(panel: PanelDefinition, rows) -> io.BytesIO:
    """通用条件查询结果导出（同样风格）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = panel.name[:28] or "查询结果"
    ws.sheet_view.showGridLines = False
    cols = panel.columns or []
    row = _title(ws, 1, len(cols), f"{panel.name} · 条件查询结果")
    row = _info(ws, row, len(cols), f"共 {len(rows)} 行　导出时间：{datetime.now():%Y-%m-%d %H:%M}")
    row = _header(ws, row, [c.get("label") for c in cols])
    for rec in rows:
        vals = [(rec.data_json or {}).get(c.get("key"), "") for c in cols]
        row = _row(ws, row, vals, aligns=LEFT)
    ws.freeze_panes = "A4"
    if cols:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{max(3, row - 1)}"
    _widths(ws, [max(10, min(30, int(c.get("width", 120)) / 8)) for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf

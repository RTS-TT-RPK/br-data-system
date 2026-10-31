import io
import json
import tempfile
from decimal import Decimal
from io import StringIO
from pathlib import Path
from unittest.mock import patch

from django.conf import settings
from django.core.management import call_command
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from panels.models import Component, IngestionBatch, PanelDefinition, Record, make_search_text
from panels.services.analysis import analyze_incoming
from panels.services.component_parser import split_components, reconcile, component_sum_percent
from panels.services.records import save_record


def _panel(key, name, columns):
    return PanelDefinition.objects.create(
        key=key, name=name, row_unit="一行 = 一个对象", columns=columns,
        form_layout=[], measurements={}, actions=[], permissions={}, version=1, status="published",
    )


def _record(panel_key, data, source="seed"):
    return Record.objects.create(
        panel_key=panel_key, panel_version=1, data_json=data,
        search_text=make_search_text(data), source_type=source,
        status=Record.STATUS_CONFIRMED,
    )


RAW_SHA_V2 = "a" * 64


def _reviewed_v2_row(row_id, *, status="ok", fields=None, components=None, image_id="测试7"):
    return {
        "row_id": row_id,
        "image_id": image_id,
        "sheet": "库存明细",
        "row_role": "data",
        "object_label": "测试库存料",
        "required_fields": [],
        "status": status,
        "raw_hash": "b" * 64,
        "fields": fields or {
            "库位": "示例库区1#", "储存区域": "示例库区1#", "编号": "P-REV-1",
            "物料名称": "测试库存料", "来源": "测试", "状态": "正常",
            "规格": "方箱", "性质": "杂原料", "数量": "3",
            "成分明细": "庚烷90% 水10%",
        },
        "components": components or [
            {"name": "庚烷", "raw_value": "90%", "value": "0.9", "unit": "%",
             "recognized": None, "from_human": False, "field": "成分明细", "revision_error": None},
            {"name": "水", "raw_value": "10%", "value": "0.1", "unit": "%",
             "recognized": None, "from_human": False, "field": "成分明细", "revision_error": None},
        ],
        "uncertain": [],
    }


def _reviewed_v2(run_id, rows, *, deliverable=True, blocked_rows=0):
    status = "deliverable" if deliverable else "draft"
    return {
        "schema_version": "reviewed/v2",
        "authority": "L2 review_state + row_overrides",
        "run_id": run_id,
        "reviewed_at": "2026-10-28T20:00:00",
        "deliverability": {
            "status": status,
            "deliverable": deliverable,
            "blocking_issues": 0 if deliverable else 1,
            "info_issues": 0,
            "blocked_rows": blocked_rows,
            "stale_decisions": 0,
            "revision_errors": 0,
            "reasons": [] if deliverable else ["还有 1 条阻塞问题未解决"],
        },
        "source": {
            "workspace": "C:/tmp/l1-fixture",
            "created": "2026-10-14 23:03:32",
            "raw_sha256": RAW_SHA_V2,
            "raw_run_id": run_id,
        },
        "stats": {"images": 1, "rows": len(rows), "sheets": 1, "issues": 0,
                  "blocking_issues": 0, "info_issues": 0, "resolved_issues": 0,
                  "blocked_rows": blocked_rows, "components": 4},
        "images": [{"image_id": "测试7", "file": "runs/20990101_001/source/测试7.png",
                    "sha256": "c" * 64, "sheet_label": "库存明细", "date_text": "",
                    "url": "/source?run_id=" + run_id + "&image_id=测试7"}],
        "sheets": [{"name": "库存明细",
                    "columns": ["库位", "储存区域", "编号", "物料名称", "来源", "状态", "规格", "性质", "数量", "成分明细"],
                    "row_count": len(rows)}],
        "rows": rows,
        "issues": [],
        "decisions": {"issue_decisions": {}, "row_overrides": {}, "archive_decisions": {}},
        "archive_plan": {},
    }


def _component(record, name, pct, unit="%", category="comp", raw=""):
    value = Decimal(str(pct)) / Decimal("100") if unit == "%" else Decimal(str(pct))
    return Component.objects.create(
        record=record, field_key="成分明细", component_std=name, component_original=name,
        numeric_value=value, unit=unit, raw_text=raw or f"{name}{pct}%", category=category,
    )


class ParserTests(TestCase):
    def test_split_and_reconcile(self):
        text = "庚烷85% 甲苯3% 水1% PH7"
        comps = split_components(text)
        self.assertEqual(len(comps), 4)
        self.assertTrue(reconcile(text, comps))
        self.assertAlmostEqual(component_sum_percent(comps), 88.0, places=6)

    def test_bad_reconcile_rejected(self):
        comps = split_components("庚烷85% 甲苯3%")
        self.assertFalse(reconcile("庚烷85% 甲苯3% 水1%", comps))


class PreviewTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        _panel("incoming_analysis", "来料分析", [
            {"key": "日期", "label": "日期", "type": "text", "width": 120},
            {"key": "位置", "label": "位置", "type": "text", "width": 120},
            {"key": "物料", "label": "物料", "type": "text", "width": 120},
            {"key": "成分明细", "label": "成分明细", "type": "components", "width": 300},
        ])
        _panel("br_inventory", "物料库存", [
            {"key": "名称", "label": "名称", "type": "text", "width": 120},
            {"key": "储位", "label": "储位", "type": "text", "width": 120},
            {"key": "储存区域", "label": "储存区域", "type": "text", "width": 120},
            {"key": "日期编号", "label": "日期编号", "type": "text", "width": 120},
            {"key": "来源", "label": "来源", "type": "text", "width": 120},
            {"key": "性质", "label": "性质", "type": "text", "width": 120},
            {"key": "成分明细", "label": "成分明细", "type": "components", "width": 300},
            {"key": "备注", "label": "备注", "type": "text", "width": 200},
            {"key": "数量/液位", "label": "数量/液位", "type": "text", "width": 100},
        ])
        _panel("br_hardware", "五金配件库", [
            {"key": "名称", "label": "名称", "type": "text", "width": 120},
            {"key": "数量", "label": "数量", "type": "text", "width": 100},
        ])
        _panel("br_material_property", "物料性质", [
            {"key": "中文名", "label": "中文名", "type": "text", "width": 120},
            {"key": "沸点(℃)", "label": "沸点", "type": "text", "width": 100},
        ])
        _panel("br_azeotrope", "共沸数据", [
            {"key": "组分A", "label": "组分A", "type": "text", "width": 120},
            {"key": "组分B", "label": "组分B", "type": "text", "width": 120},
            {"key": "共沸点(℃)", "label": "共沸点", "type": "text", "width": 100},
            {"key": "备注", "label": "备注", "type": "text", "width": 200},
        ])

        _record("br_material_property", {"中文名": "正庚烷", "沸点(℃)": "98.4"})
        _record("br_material_property", {"中文名": "甲苯", "沸点(℃)": "111"})
        _record("br_azeotrope", {"组分A": "正庚烷", "组分B": "甲苯", "共沸点(℃)": "无恒定共沸", "备注": "不共沸"})

        incoming = _record("incoming_analysis", {
            "日期": "2026-10-27", "位置": "示例库区1#", "物料": "示例混料2",
            "成分明细": "庚烷85% 甲苯3% 水1%",
        })
        _component(incoming, "庚烷", 85, category="comp", raw="庚烷85%")
        _component(incoming, "甲苯", 3, category="comp", raw="甲苯3%")
        _component(incoming, "水", 1, category="water", raw="水1%")

        inv = _record("br_inventory", {"名称": "示例混料2", "储位": "示例库区1#", "储存区域": "示例库区1#", "成分明细": "庚烷90% 甲苯3%"})
        _component(inv, "庚烷", 90, category="comp", raw="庚烷90%")
        _component(inv, "甲苯", 3, category="comp", raw="甲苯3%")

        hw = _record("br_hardware", {"名称": "螺丝", "数量": "10"})
        _component(hw, "测试", 0)  # 只为覆盖模型

    def test_pages(self):
        for url in ["/", "/p/incoming_analysis/", "/p/br_inventory/", "/analysis/incoming/", "/data/query/", "/query/"]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)

    def test_ingestion_api_and_approve(self):
        payload = {
            "panel_key": "incoming_analysis",
            "source": {"type": "test", "run_id": "T-001"},
            "rows": [{
                "row_key": "R1",
                "fields": {"日期": "2026-10-27", "位置": "示例库区2#", "物料": "测试料", "成分明细": "庚烷80% 甲苯2%"},
                "components": [{"name": "庚烷", "raw_value": "80%"}, {"name": "甲苯", "raw_value": "2%"}],
            }],
        }
        r = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(r.status_code, 201)
        batch = IngestionBatch.objects.get(pk=r.json()["batch_id"])
        self.assertEqual(batch.rows.count(), 1)
        self.client.post(f"/ingestions/{batch.pk}/approve/", {})
        batch.refresh_from_db()
        self.assertEqual(batch.status, "approved")
        self.assertTrue(any((r.data_json or {}).get("物料") == "测试料" for r in Record.objects.filter(panel_key="incoming_analysis")))


    def test_import_reviewed_creates_draft_then_approve(self):
        reviewed = _reviewed_v2("T-REV-1", [_reviewed_v2_row("R-REV-1")])
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "reviewed.json"
            p.write_text(json.dumps(reviewed, ensure_ascii=False), encoding="utf-8")
            before = Record.objects.filter(panel_key="br_inventory").count()
            out = StringIO()
            call_command(
                "import_reviewed",
                reviewed=str(p),
                mapping=str(Path(settings.BASE_DIR) / "seeds" / "reviewed_maps.json"),
                stdout=out,
            )
            self.assertIn("待审核", out.getvalue())
            self.assertIn("可导入 1 行", out.getvalue())
            batch = IngestionBatch.objects.get(source_run_id="T-REV-1")
            self.assertEqual(batch.status, "pending")
            self.assertEqual(batch.rows.count(), 1)
            self.assertEqual(Record.objects.filter(panel_key="br_inventory").count(), before)
            self.client.post(f"/ingestions/{batch.pk}/approve/", {})
            batch.refresh_from_db()
            self.assertEqual(batch.status, "approved")
            rec = Record.objects.filter(panel_key="br_inventory", source_run_id="T-REV-1").latest("id")
            self.assertEqual(rec.data_json["名称"], "测试库存料")
            self.assertEqual(rec.data_json["储位"], "示例库区1#")
            self.assertEqual(rec.components.count(), 2)
            self.assertEqual(rec.components.get(component_std="庚烷").unit, "%")
            self.assertAlmostEqual(float(rec.components.get(component_std="庚烷").numeric_value), 0.9, places=6)

    def test_import_reviewed_is_idempotent(self):
        reviewed = _reviewed_v2("T-REV-2", [
            _reviewed_v2_row("R-REV-2", fields={
                "库位": "示例库区2#", "储存区域": "示例库区2#", "编号": "P-REV-2",
                "物料名称": "重复料", "来源": "测试", "状态": "正常", "规格": "方箱",
                "性质": "杂原料", "数量": "1", "成分明细": "庚烷80%",
            }, components=[
                {"name": "庚烷", "raw_value": "80%", "value": "0.8", "unit": "%"},
            ]),
        ])
        with tempfile.TemporaryDirectory() as td:
            p = Path(td) / "reviewed.json"
            p.write_text(json.dumps(reviewed, ensure_ascii=False), encoding="utf-8")
            kwargs = dict(reviewed=str(p), mapping=str(Path(settings.BASE_DIR) / "seeds" / "reviewed_maps.json"))
            call_command("import_reviewed", **kwargs)
            call_command("import_reviewed", **kwargs)
            self.assertEqual(IngestionBatch.objects.filter(source_run_id="T-REV-2").count(), 1)
            self.assertEqual(IngestionBatch.objects.get(source_run_id="T-REV-2").rows.count(), 1)

    def test_combo_query(self):
        r = self.client.get("/query/", {
            "panel_key": "br_inventory", "material": "示例混料2",
            "cond_name_1": "庚烷", "cond_op_1": "gte", "cond_val_1": "80",
            "cond_name_2": "甲苯", "cond_op_2": "lte", "cond_val_2": "5",
        })
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "示例混料2")
        self.assertContains(r, "庚烷")
        self.assertContains(r, "条件核对")

    def test_combo_query_alias(self):
        r = self.client.get("/query/", {
            "panel_key": "br_inventory", "material": "示例混料2",
            "cond_name_1": "正庚烷", "cond_op_1": "gte", "cond_val_1": "80",
        })
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "示例混料2")

    def test_data_query(self):
        r = self.client.get("/data/query/", {
            "panel_key": "br_inventory", "field_1": "名称", "op_1": "contains", "val_1": "示例混料",
        })
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "示例混料2")

    def test_report_xlsx(self):
        r = self.client.get("/export/summary.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        self.assertIn("综合汇总", wb.sheetnames)
        self.assertIn("物料成分明细", wb.sheetnames)
        wb.close()


    def test_manual_crud(self):
        r = self.client.post("/p/incoming_analysis/new/", {
            "日期": "2026-10-28", "位置": "示例库区3#", "物料": "手动料",
            "成分明细": "庚烷80% 甲苯2%",
        })
        self.assertEqual(r.status_code, 302)
        rec = Record.objects.filter(panel_key="incoming_analysis").order_by("-id").first()
        self.assertEqual(rec.data_json.get("物料"), "手动料")
        self.assertEqual(rec.components.count(), 2)
        r = self.client.post(f"/p/incoming_analysis/{rec.pk}/edit/", {
            "日期": "2026-10-29", "位置": "示例库区3#", "物料": "手动料2",
            "成分明细": "庚烷81% 甲苯2%",
        })
        self.assertEqual(r.status_code, 302)
        rec.refresh_from_db()
        self.assertEqual(rec.data_json.get("物料"), "手动料2")
        self.client.post(f"/p/incoming_analysis/{rec.pk}/delete/")
        rec.refresh_from_db()
        self.assertEqual(rec.status, Record.STATUS_DELETED)

    def test_csv_export(self):
        r = self.client.get("/p/br_inventory/export.csv")
        self.assertEqual(r.status_code, 200)
        self.assertIn("text/csv", r["Content-Type"])
        self.assertTrue(r.content.startswith("\ufeff".encode("utf-8")))

    def test_component_query(self):
        r = self.client.get("/components/", {"name": "庚烷"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "庚烷")

    def test_data_query_export(self):
        r = self.client.get("/data/query.xlsx", {"panel_key": "br_inventory", "field_1": "名称", "op_1": "contains", "val_1": "庚烷"})
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        self.assertIn("物料库存", wb.sheetnames)
        wb.close()

    def test_api_ingestion_detail(self):
        payload = {"panel_key": "incoming_analysis", "rows": [{"fields": {"物料": "X"}}]}
        r = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json")
        batch_id = r.json()["batch_id"]
        d = self.client.get(f"/api/ingestions/{batch_id}")
        self.assertEqual(d.status_code, 200)
        self.assertEqual(d.json()["rows"], 1)

    @override_settings(API_TOKEN="test-" + "secret")
    def test_api_token_required(self):
        token = "test-" + "secret"
        payload = {"panel_key": "incoming_analysis", "rows": [{"fields": {"物料": "X"}}]}
        r = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(r.status_code, 401)
        r = self.client.post(
            "/api/ingestions", data=json.dumps(payload), content_type="application/json",
            HTTP_X_API_KEY=token,
        )
        self.assertEqual(r.status_code, 201)


    def test_golden_analysis_values(self):
        rows = analyze_incoming()
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["main_name"], "庚烷")
        self.assertAlmostEqual(row["main_pct"], 85.0, places=6)
        self.assertAlmostEqual(row["main_bp"], 98.4, places=6)
        self.assertAlmostEqual(row["min_gap"], 1.6, places=6)
        self.assertAlmostEqual(row["max_gap"], 12.6, places=6)
        self.assertEqual(row["azeo_label"], "部分未收录")
        by_name = {o["name"]: o for o in row["others"]}
        self.assertAlmostEqual(by_name["甲苯"]["gap"], 12.6, places=6)
        self.assertAlmostEqual(by_name["水"]["gap"], 1.6, places=6)

    def test_transaction_rollback_on_component_failure(self):
        panel = PanelDefinition.objects.get(key="incoming_analysis")
        before = Record.objects.count()
        with patch("panels.services.records.Component.objects.bulk_create", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                save_record(panel, {"日期": "2026-10-28", "位置": "示例库区3#", "物料": "事务料", "成分明细": "庚烷85%"})
        self.assertEqual(Record.objects.count(), before)
        self.assertFalse(any((r.data_json or {}).get("物料") == "事务料" for r in Record.objects.all()))

    def test_edit_rollback_keeps_old_components(self):
        panel = PanelDefinition.objects.get(key="incoming_analysis")
        rec = _record("incoming_analysis", {"日期": "2026-10-28", "物料": "回滚料", "成分明细": "庚烷85% 甲苯3%"})
        _component(rec, "庚烷", 85, raw="庚烷85%")
        _component(rec, "甲苯", 3, raw="甲苯3%")
        before = rec.components.count()
        with patch("panels.services.records.Component.objects.bulk_create", side_effect=RuntimeError("boom")):
            with self.assertRaises(RuntimeError):
                save_record(panel, {"日期": "2026-10-28", "物料": "回滚料2", "成分明细": "庚烷80% 甲苯4%"}, record=rec)
        rec.refresh_from_db()
        self.assertEqual(rec.data_json.get("物料"), "回滚料")
        self.assertEqual(rec.components.count(), before)

    def test_manual_page_rejects_unparsable_components(self):
        panel = PanelDefinition.objects.get(key="incoming_analysis")
        before = Record.objects.count()
        with patch("panels.services.records.split_components", return_value=[]):
            r = self.client.post("/p/incoming_analysis/new/", {
                "日期": "2026-10-28", "位置": "示例库区3#", "物料": "坏料", "成分明细": "庚烷85%",
            })
        self.assertEqual(r.status_code, 200)
        self.assertEqual(Record.objects.count(), before)

    def test_duplicate_post_is_idempotent(self):
        payload = {
            "panel_key": "incoming_analysis",
            "source": {"type": "test", "run_id": "IDEMP-001"},
            "rows": [{"row_key": "R1", "fields": {"物料": "幂等料", "成分明细": "庚烷85% 甲苯2%"}}],
        }
        r1 = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json")
        r2 = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(r1.status_code, 201)
        self.assertEqual(r2.status_code, 200)
        self.assertTrue(r2.json()["duplicate"])
        self.assertEqual(r1.json()["batch_id"], r2.json()["batch_id"])
        self.assertEqual(IngestionBatch.objects.filter(source_run_id="IDEMP-001").count(), 1)

    def test_duplicate_row_key_rejected(self):
        payload = {
            "panel_key": "incoming_analysis",
            "rows": [
                {"row_key": "R1", "fields": {"物料": "A"}},
                {"row_key": "R1", "fields": {"物料": "B"}},
            ],
        }
        r = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("row_key", r.json()["error"])

    def test_approve_duplicate_and_reject_guard(self):
        payload = {
            "panel_key": "incoming_analysis",
            "source": {"type": "test", "run_id": "APPROVE-001"},
            "rows": [{"row_key": "R1", "fields": {"物料": "确认料", "成分明细": "庚烷85%"}}],
        }
        batch_id = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json").json()["batch_id"]
        before = Record.objects.count()
        self.client.post(f"/ingestions/{batch_id}/approve/", {})
        self.assertEqual(Record.objects.count(), before + 1)
        self.client.post(f"/ingestions/{batch_id}/approve/", {})
        self.assertEqual(Record.objects.count(), before + 1)
        self.client.post(f"/ingestions/{batch_id}/reject/", {})
        batch = IngestionBatch.objects.get(pk=batch_id)
        self.assertEqual(batch.status, "approved")

    def test_consistency_command(self):
        out = StringIO()
        call_command("check_consistency", stdout=out)
        self.assertIn("一致性检查通过", out.getvalue())

    def test_report_key_cells(self):
        r = self.client.get("/export/summary.xlsx")
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["综合汇总"]
        self.assertEqual(ws["A1"].value, "示例企业物料综合汇总表")
        self.assertEqual(ws["A4"].value, "物料种类")
        self.assertEqual(ws["A4"].fill.fgColor.rgb, "00D6E4F0")
        wb.close()


    def test_simulation_is_idempotent(self):
        before = IngestionBatch.objects.count()
        r1 = self.client.post("/ingestions/simulate/")
        r2 = self.client.post("/ingestions/simulate/")
        self.assertEqual(r1.status_code, 302)
        self.assertEqual(r2.status_code, 302)
        self.assertEqual(IngestionBatch.objects.count(), before + 1)

    def test_approval_blocks_missing_components(self):
        payload = {
            "panel_key": "incoming_analysis",
            "source": {"type": "test", "run_id": "MISSING-001"},
            "rows": [{"row_key": "R1", "fields": {"物料": "缺成分"}}],
        }
        batch_id = self.client.post("/api/ingestions", data=json.dumps(payload), content_type="application/json").json()["batch_id"]
        before = Record.objects.count()
        self.client.post(f"/ingestions/{batch_id}/approve/", {})
        self.assertEqual(Record.objects.count(), before)
        row = IngestionBatch.objects.get(pk=batch_id).rows.get()
        self.assertTrue(any("缺少成分原文" in x for x in row.issues))

    def test_agent_apis(self):
        self.assertEqual(self.client.get("/api/health").status_code, 200)
        r = self.client.get("/api/records", {"panel_key": "br_inventory"})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["count"], 1)
        r = self.client.get("/api/analysis/incoming")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["count"], 1)
